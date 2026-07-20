"""인사이트(Insights) — 시스템 로그 · 사용량 통계 · DB 구조 · DB 도구 열기 (관리자 콘솔 §9·10·6).

- logs: AuditLog 를 단일 소스로 event/network/dicom 분류(액션 프리픽스 매핑) + 날짜/병원/검색 필터,
  같은 쿼리의 /logs.csv 는 CSV 다운로드(BOM 포함 — 엑셀 한글 안전).
- stats: Study 집계(병원/장비/진료과/판독상태 × 기간 × 병원 스코프) — 병원별 탭(hid)과 공용.
  같은 쿼리의 /stats.xlsx 는 openpyxl 기반 Excel 다운로드(헤더 볼드·열 너비·숫자 서식).
- db-schema: SQLAlchemy inspector 로 read-only 구조 확인(요구 6 "DB 프로그램 열기"의 내장 뷰).
- db-tool-open: 설정 server.dbtool.path 의 외부 DB 도구를 서버측에서 분리 실행(관리자 전용).

전부 관리자 전용(admin_user). 라우터 등록(main.py)은 B1 레인 담당.
"""
from __future__ import annotations

import csv
import io
import json
import os
import subprocess
from datetime import datetime, timedelta, timezone
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import and_, func, or_, select, text
from sqlalchemy import inspect as sa_inspect
from sqlalchemy.orm import Session

from app.api.deps import admin_user
from app.db import get_db
from app.models import Account, AuditLog, Hospital, Study
from app.services.settings_service import get_setting

router = APIRouter(prefix="/api/insights", tags=["insights"])


# ════════════════════════════════ 시스템 로그 (요구 9) ════════════════════════════════
# AuditLog 액션 → 로그 타입 분류. 별도 network/dicom 로그 저장소가 없어(감사 로그 단일 소스)
# 액션 이름의 집합·프리픽스 매핑으로 분류한다. 새 액션은 프리픽스 규칙으로 자동 흡수.
_DICOM_ACTIONS = {
    "import_dicom", "send_gsps", "send_kos", "mpps_update", "dicom_nodes_apply",
    "modality_apply", "report_send_sr", "mwl_export", "scp_config",
}
_DICOM_PREFIXES = ("dicom", "mpps", "send_", "mwl", "scp", "modality_apply", "import_dicom")
_NETWORK_ACTIONS = {"login", "login_failed", "login_disabled", "client_enter", "heartbeat"}
_NETWORK_PREFIXES = ("login", "net_", "client_", "heartbeat")

LOG_TYPES = ("event", "network", "dicom")
_FETCH_CAP = 5000  # 타입 분류는 파이썬에서 하므로 과대 스캔 방지 상한


def _classify(action: str) -> str:
    a = (action or "").lower()
    if a in _DICOM_ACTIONS or a.startswith(_DICOM_PREFIXES):
        return "dicom"
    if a in _NETWORK_ACTIONS or a.startswith(_NETWORK_PREFIXES):
        return "network"
    return "event"


def _parse_dt(value: str, *, end: bool = False) -> datetime | None:
    """'YYYY-MM-DD'(또는 ISO datetime) → aware UTC. end=True 이고 날짜만이면 익일 0시(미만 비교용)."""
    s = (value or "").strip()
    if not s:
        return None
    try:
        dt = datetime.fromisoformat(s)
    except ValueError:
        raise HTTPException(status_code=400, detail=f"날짜 형식 오류: {s} (YYYY-MM-DD)")
    date_only = len(s) <= 10
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    if end and date_only:
        dt = dt + timedelta(days=1)
    return dt


def _query_logs(db: Session, *, type_: str, date_from: str, date_to: str,
                hid: int | None, q: str, limit: int) -> list[dict]:
    if type_ and type_ not in LOG_TYPES:
        raise HTTPException(status_code=400, detail="type은 event|network|dicom")
    limit = max(1, min(limit, 2000))
    conds = []
    dt_from = _parse_dt(date_from)
    dt_to = _parse_dt(date_to, end=True)
    if dt_from is not None:
        conds.append(AuditLog.created_at >= dt_from)
    if dt_to is not None:
        conds.append(AuditLog.created_at < dt_to)
    if hid is not None:
        # 행위자 소속 병원 또는 병원을 대상으로 한 이벤트
        conds.append(or_(
            Account.hospital_id == hid,
            and_(AuditLog.target_type == "hospital", AuditLog.target_id == str(hid)),
        ))
    if q.strip():
        like = f"%{q.strip()}%"
        conds.append(or_(
            AuditLog.action.ilike(like), AuditLog.target_type.ilike(like),
            AuditLog.target_id.ilike(like), Account.username.ilike(like),
        ))
    stmt = (
        select(AuditLog, Account.username, Account.hospital_id)
        .join(Account, AuditLog.account_id == Account.id, isouter=True)
        .where(*conds)
        .order_by(AuditLog.id.desc())
        .limit(_FETCH_CAP)
    )
    items: list[dict] = []
    for log, username, acc_hid in db.execute(stmt):
        t = _classify(log.action)
        if type_ and t != type_:
            continue
        hospital_id = acc_hid
        if hospital_id is None and log.target_type == "hospital":
            try:
                hospital_id = int(log.target_id)
            except (TypeError, ValueError):
                hospital_id = None
        items.append({
            "ts": log.created_at.isoformat() if log.created_at else "",
            "type": t,
            "actor": username or "",
            "hospital_id": hospital_id,
            "action": log.action,
            "detail": {"target_type": log.target_type, "target_id": log.target_id,
                       **(log.detail or {})},
        })
        if len(items) >= limit:
            break
    return items


@router.get("/logs")
def logs(type: str = "", date_from: str = "", date_to: str = "", hid: int | None = None,
         q: str = "", limit: int = 200,
         db: Session = Depends(get_db), user: dict = Depends(admin_user)):
    """시스템 로그 — event/network/dicom 분류 + 날짜·병원·검색 필터."""
    return {"items": _query_logs(db, type_=type, date_from=date_from, date_to=date_to,
                                 hid=hid, q=q, limit=limit)}


@router.get("/logs.csv")
def logs_csv(type: str = "", date_from: str = "", date_to: str = "", hid: int | None = None,
             q: str = "", limit: int = 2000,
             db: Session = Depends(get_db), user: dict = Depends(admin_user)):
    """같은 쿼리의 CSV 다운로드 — BOM(utf-8-sig) 포함으로 엑셀 한글 안전."""
    items = _query_logs(db, type_=type, date_from=date_from, date_to=date_to,
                        hid=hid, q=q, limit=limit)

    def _rows():
        buf = io.StringIO()
        w = csv.writer(buf)
        yield "﻿"  # UTF-8 BOM (엑셀 한글 인식)
        w.writerow(["ts", "type", "actor", "hospital_id", "action", "detail"])
        yield buf.getvalue()
        for it in items:
            buf.seek(0)
            buf.truncate(0)
            w.writerow([it["ts"], it["type"], it["actor"],
                        it["hospital_id"] if it["hospital_id"] is not None else "",
                        it["action"], json.dumps(it["detail"], ensure_ascii=False)])
            yield buf.getvalue()

    fname = f"system-logs-{datetime.now(timezone.utc).strftime('%Y%m%d')}.csv"
    return StreamingResponse(
        _rows(), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# ════════════════════════════════ 사용량 통계 (요구 10) ════════════════════════════════
STAT_GROUPS = ("hospital", "modality", "department", "report_status")
# 그룹 → 한국어 라벨 (Excel 시트명 등 사용자 노출용)
STAT_GROUP_LABELS = {"hospital": "병원별", "modality": "장비별",
                     "department": "진료과별", "report_status": "판독상태"}


def _norm_date8(value: str) -> str:
    """'YYYY-MM-DD' | 'YYYYMMDD' → 'YYYYMMDD' (Study.study_date 비교용)."""
    return (value or "").replace("-", "").strip()


def _query_stats(db: Session, *, group: str, date_from: str, date_to: str,
                 hid: int | None) -> dict:
    """통계 집계 본체 — /stats(JSON)와 /stats.xlsx(Excel)가 동일 쿼리를 공유한다."""
    if group not in STAT_GROUPS:
        raise HTTPException(status_code=400, detail="group은 hospital|modality|department|report_status")
    conds = []
    df, dt = _norm_date8(date_from), _norm_date8(date_to)
    if df:
        conds.append(Study.study_date >= df)
    if dt:
        conds.append(Study.study_date <= dt)
    if hid is not None:
        conds.append(Study.hospital_id == hid)

    finalized = Study.status == "finalized"  # 판독완료 = 확정(finalized) 상태
    if group == "report_status":
        total = db.execute(select(func.count()).select_from(Study).where(*conds)).scalar() or 0
        fin = db.execute(
            select(func.count()).select_from(Study).where(*conds, finalized)
        ).scalar() or 0
        rows = [
            {"key": "finalized", "label": "판독완료", "studies": fin, "reports": fin, "unreported": 0},
            {"key": "unreported", "label": "미판독", "studies": total - fin, "reports": 0,
             "unreported": total - fin},
        ]
        return {"group": group, "rows": rows}

    col = {"hospital": Study.hospital_id, "modality": Study.modality,
           "department": Study.department}[group]
    totals = dict(db.execute(select(col, func.count()).where(*conds).group_by(col)).all())
    finals = dict(db.execute(
        select(col, func.count()).where(*conds, finalized).group_by(col)
    ).all())
    names = {h.id: (h.name or h.code) for h in db.execute(select(Hospital)).scalars().all()} \
        if group == "hospital" else {}
    rows = []
    for k in sorted(totals, key=lambda x: (x is None, str(x))):
        n = totals[k]
        f = finals.get(k, 0)
        if group == "hospital":
            key = str(k) if k is not None else ""
            label = names.get(k, "(미배정)") if k is not None else "(미배정)"
        else:
            key = k or ""
            label = k or "(미지정)"
        rows.append({"key": key, "label": label, "studies": n, "reports": f, "unreported": n - f})
    return {"group": group, "rows": rows}


@router.get("/stats")
def stats(group: str = "hospital", date_from: str = "", date_to: str = "",
          hid: int | None = None,
          db: Session = Depends(get_db), user: dict = Depends(admin_user)):
    """사용량 통계 — 병원/장비/진료과/판독상태별 검사·판독 집계 (기간·병원 스코프)."""
    return _query_stats(db, group=group, date_from=date_from, date_to=date_to, hid=hid)


@router.get("/stats.xlsx")
def stats_xlsx(group: str = "hospital", date_from: str = "", date_to: str = "",
               hid: int | None = None,
               db: Session = Depends(get_db), user: dict = Depends(admin_user)):
    """같은 쿼리의 Excel(.xlsx) 다운로드 — 헤더 볼드·열 너비·숫자 서식·시트명=그룹 라벨."""
    data = _query_stats(db, group=group, date_from=date_from, date_to=date_to, hid=hid)

    from openpyxl import Workbook  # 지연 임포트 — /stats(JSON) 경로는 openpyxl 없이도 동작
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = STAT_GROUP_LABELS.get(group, group)
    ws.append(["구분", "검사 수", "판독", "미판독"])
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
    total = {"studies": 0, "reports": 0, "unreported": 0}
    for r in data["rows"]:
        ws.append([r["label"] or r["key"], r["studies"], r["reports"], r["unreported"]])
        for k in total:
            total[k] += r[k]
    if data["rows"]:  # 합계 행(볼드) — 빈 결과면 생략
        ws.append(["합계", total["studies"], total["reports"], total["unreported"]])
        for cell in ws[ws.max_row]:
            cell.font = bold
    for col, width in (("A", 26), ("B", 12), ("C", 12), ("D", 12)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=4):
        for cell in row:
            cell.number_format = "#,##0"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"usage-stats-{group}-{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# ════════════════════════════════ DB 구조 확인 (요구 6) ════════════════════════════════
@router.get("/db-schema")
def db_schema(db: Session = Depends(get_db), user: dict = Depends(admin_user)):
    """read-only 스키마 조회 — 테이블·컬럼·행수 (외부 DB 도구 없이도 구조 확인)."""
    bind = db.get_bind()
    insp = sa_inspect(bind)
    tables = []
    for name in sorted(insp.get_table_names()):
        cols = [{"name": c["name"], "type": str(c["type"])} for c in insp.get_columns(name)]
        try:
            # 테이블명은 inspector 가 반환한 실제 스키마 이름만 사용(임의 입력 아님)
            n_rows = db.execute(text(f'SELECT COUNT(*) FROM "{name}"')).scalar() or 0
        except Exception:  # noqa: BLE001 — 뷰/권한 등으로 카운트 불가 시 -1
            n_rows = -1
        tables.append({"name": name, "rows": n_rows, "columns": cols})
    return {"tables": tables}


# ════════════════════════════════ DB 도구 실행 (요구 6) ════════════════════════════════
DBTOOL_KEY = "server.dbtool"


@router.post("/db-tool-open")
def db_tool_open(db: Session = Depends(get_db), user: dict = Depends(admin_user)):
    """설정 server.dbtool.path 의 외부 DB 프로그램을 서버측에서 분리 실행.

    보안: 관리자 전용 · 관리자가 설정 화면에서 저장한 경로만 실행 · 인자 주입 없음(단일 인자 배열).
    """
    cfg = get_setting(db, DBTOOL_KEY, default={}) or {}
    path = str(cfg.get("path", "")).strip()
    if not path:
        raise HTTPException(status_code=400,
                            detail="DB 도구 경로가 설정되지 않았습니다 (설정 키 server.dbtool)")
    p = Path(path)
    if not p.is_file():
        raise HTTPException(status_code=400, detail=f"실행 파일을 찾을 수 없습니다: {path}")
    creationflags = 0
    if os.name == "nt":  # 서버 프로세스와 분리(콘솔/프로세스 그룹 독립)
        creationflags = subprocess.DETACHED_PROCESS | subprocess.CREATE_NEW_PROCESS_GROUP
    try:
        subprocess.Popen([str(p)], cwd=str(p.parent), close_fds=True,
                         creationflags=creationflags)
    except OSError as e:
        raise HTTPException(status_code=500, detail=f"실행 실패: {str(e)[:80]}")
    db.add(AuditLog(account_id=user.get("uid"), action="db_tool_open",
                    target_type="server", target_id=p.name, detail={"path": str(p)}))
    db.commit()
    return {"ok": True, "detail": f"실행됨: {p.name}"}


@router.get("/db-tool-detect")
def db_tool_detect(user: dict = Depends(admin_user)):
    """서버 PC 에 설치된 흔한 DB 도구 자동 탐지 — 경로 미설정 시 기본값 제안용.

    잘 알려진 설치 경로만 확인(디스크 전체 스캔 없음). 첫 항목이 기본 제안값.
    """
    import glob as _glob

    candidates: list[tuple[str, str]] = []  # (label, glob pattern)
    for base in (r"C:\Program Files", r"C:\Program Files (x86)"):
        candidates += [
            ("pgAdmin 4", rf"{base}\pgAdmin 4\runtime\pgAdmin4.exe"),
            ("pgAdmin 4", rf"{base}\pgAdmin 4\v*\runtime\pgAdmin4.exe"),
            ("DBeaver", rf"{base}\DBeaver\dbeaver.exe"),
            ("HeidiSQL", rf"{base}\HeidiSQL\heidisql.exe"),
            ("DB Browser for SQLite", rf"{base}\DB Browser for SQLite\DB Browser for SQLite.exe"),
            ("TablePlus", rf"{base}\TablePlus\TablePlus.exe"),
        ]
    local = os.environ.get("LOCALAPPDATA", "")
    if local:
        candidates += [
            ("DBeaver", rf"{local}\DBeaver\dbeaver.exe"),
            ("HeidiSQL", rf"{local}\Programs\HeidiSQL\heidisql.exe"),
        ]
    items: list[dict] = []
    seen: set[str] = set()
    for label, pat in candidates:
        for hit in _glob.glob(pat):
            if hit not in seen and Path(hit).is_file():
                seen.add(hit)
                items.append({"label": label, "path": hit})
    return {"items": items}
